import re

import openpyxl
import pytest

# ============================================================
# LƯU KẾT QUẢ TEST
# ============================================================
RESULTS = {}          # {"TC-005": "Pass" | "Fail"}
ACTUAL_RESULTS = {}   # {"TC-005": ["dòng 1", "dòng 2", ...]}
PENDING = {}          # dành cho save_actual() kiểu cũ: {"TC-005": ["..."]}

XLSX = r"D:\Excel\Test_docs_completed1_TC_trung_Word.xlsx"

# Test name dạng: test_tc005_xxx | test_tc037_038_039_xxx (nhiều TC trong 1 test)
TC_PATTERN = re.compile(r"test_tc(\d{3}(?:_\d{3})*)")


# ============================================================
# FIXTURE LOGIN
# ============================================================
@pytest.fixture
def login(page):
    def _login(username, password):
        page.goto("/")
        page.get_by_text("Đăng nhập").first.click()
        page.fill("#loginUsername", username)
        page.fill("#loginPass", password)
        page.click("#formLogin button[type='submit']")

    return _login


# ============================================================
# GHI ACTUAL RESULT
# ============================================================
@pytest.fixture
def actual(record_property):
    """
    Dùng trong test:   actual("HTTP 403, không trả dữ liệu")
    Tự biết test thuộc TC nào (lấy từ tên hàm). Gọi nhiều lần được -> nối nhiều dòng.
    Nên gọi TRƯỚC các câu assert để test fail vẫn có số liệu thực tế.
    """
    def _actual(text):
        record_property("actual", str(text))

    return _actual


def save_actual(tc_id, actual_result):
    """Cách cũ (vẫn dùng được): save_actual("TC-005", "...")"""
    PENDING.setdefault(tc_id, []).append(str(actual_result))


# ============================================================
# LẤY TC ID + PASS/FAIL + ACTUAL
# ============================================================
def _short_error(report):
    lr = report.longrepr
    crash = getattr(lr, "reprcrash", None)
    msg = crash.message if crash else str(lr)
    return msg.strip()[:600]


def pytest_runtest_logreport(report):
    if report.when != "call":
        return

    m = TC_PATTERN.search(report.nodeid)
    if not m:
        return

    tc_ids = [f"TC-{n}" for n in m.group(1).split("_")]

    # nhãn tham số, vd test_tc018_role_hien_thi[admin] -> "admin"
    lab = re.search(r"\[(.+)\]$", report.nodeid)
    label = f"[{lab.group(1)}] " if lab else ""

    notes = [v for k, v in report.user_properties if k == "actual"]

    for tc_id in tc_ids:
        lines = list(notes) + PENDING.pop(tc_id, [])

        if report.passed:
            text = "; ".join(lines) or "Test thực thi thành công."
        else:
            text = "; ".join(lines)
            text = (text + " | " if text else "") + "Test thất bại: " + _short_error(report)

        # Một TC có nhiều test/tham số: chỉ cần 1 cái Fail là TC Fail
        if RESULTS.get(tc_id) == "Fail" or not report.passed:
            RESULTS[tc_id] = "Fail"
        else:
            RESULTS[tc_id] = "Pass"

        ACTUAL_RESULTS.setdefault(tc_id, []).append(label + text)

        print(f"\n{tc_id}: {RESULTS[tc_id]}")
        print(f"Actual Result: {label + text}")


# ============================================================
# GHI KẾT QUẢ VÀO EXCEL
# ============================================================
def pytest_sessionfinish(session):
    if not RESULTS:
        return

    try:
        wb = openpyxl.load_workbook(XLSX)
        ws = wb["TestCase"]

        for r in range(2, ws.max_row + 1):
            tid = ws.cell(r, 1).value      # Cột A = Test Case ID
            if not tid:
                continue
            tid = str(tid).strip()

            if tid in ACTUAL_RESULTS:      # Cột G = Actual Result
                ws.cell(r, 7).value = "\n".join(ACTUAL_RESULTS[tid])
            if tid in RESULTS:             # Cột H = Pass/Fail
                ws.cell(r, 8).value = RESULTS[tid]

        wb.save(XLSX)

    except PermissionError:
        print(f"\n[WARNING] Không thể ghi Excel: {XLSX}\nHãy đóng file Excel rồi chạy lại.")
        return
    except Exception as e:
        print(f"\n[ERROR] Không thể cập nhật Excel:\n{e}")
        return

    print("\n===== KẾT QUẢ TEST =====")
    for tc_id in sorted(RESULTS):
        print(f"{tc_id}: {RESULTS[tc_id]}")
        print("Actual Result: " + " | ".join(ACTUAL_RESULTS.get(tc_id, [])))
    print(f"\nĐã cập nhật Excel: {XLSX}")